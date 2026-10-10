import openpyxl
import pandas as pd
import numpy as np
import re
import os

def norm_cod(c):
    """Normaliza el código de EPS a un entero o identificador limpio."""
    if c is None or pd.isna(c):
        return None
    s = str(c).strip().split('-')[0]
    nums = re.sub(r'\D', '', s)
    return int(nums) if nums else None

def norm_periodo(p):
    """Extrae el número romano del periodo regulatorio (ej: 'III PR' -> 'III')."""
    if not p or pd.isna(p):
        return ''
    m = re.search(r'\b(I{1,3}|IV|V)\b', str(p).upper())
    return m.group(1) if m else str(p).strip().upper()

def norm_anio(a):
    """Normaliza la etiqueta del año regulatorio."""
    if not a or pd.isna(a):
        return ''
    s = str(a).strip().upper()
    if 'BASE' in s:
        return 'Año Base'
    if 'ACUM' in s:
        return 'Acumulado'
    m_pt = re.search(r'PT\s*(\d+)', s)
    if m_pt:
        return f"PT {m_pt.group(1)}"
    m = re.search(r'([1-9])', s)
    return f"Año {m.group(1)}" if m else s

def limpiar_link(target):
    """Normaliza enlaces de SharePoint/OneDrive a URLs web absolutas y navegables."""
    if not target:
        return None
    s = str(target).strip()
    if s.startswith('http'):
        return s
    if ':b:/g/personal/' in s:
        idx = s.find(':b:/g/personal/')
        return 'https://sunassgobpe-my.sharepoint.com/' + s[idx:]
    return s

def limpiar_valor(valor, unidad_medida, es_ici=False):
    """Limpia y castea valores de metas, ejecuciones e ICIs."""
    if pd.isna(valor) or str(valor).strip() in ('-', '', 'None', 'nan', 'NaN'):
        return None
        
    # Corrección de fechas accidentales de Excel (ej: 14.8 convertido a 14/08)
    if isinstance(valor, pd.Timestamp) or type(valor).__name__ == 'datetime':
        return float(f"{valor.day}.{valor.month}")
        
    val_str = str(valor).strip()
    if es_ici or (unidad_medida and str(unidad_medida).strip() == '%'):
        if val_str.endswith('%'):
            val_str = val_str[:-1].strip()
        try:
            val_float = float(val_str)
            if 0 <= val_float <= 1.5:
                return round(val_float * 100, 2)
            else:
                return round(val_float, 2)
        except (ValueError, TypeError):
            pass
            
    try:
        val_float = float(valor)
        return round(val_float, 2)
    except (ValueError, TypeError):
        pass
        
    return str(valor).strip()

def cargar_informes_totales(file_path):
    """Carga y mapea todos los informes e hipervínculos de la hoja 'INFORMES TOTALES'."""
    print("Cargando hoja 'INFORMES TOTALES' e hipervínculos...")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb['INFORMES TOTALES']
    
    dict_inf = {}
    dict_eps_meta = {} # Metadatos a nivel de EPS (Departamento, Tamaño)
    
    for r in range(7, ws.max_row + 1):
        c_cod = ws.cell(r, 1).value
        c_dep = ws.cell(r, 2).value
        c_nom = ws.cell(r, 3).value
        c_per = ws.cell(r, 4).value
        c_eval = ws.cell(r, 14).value
        c_est = ws.cell(r, 15).value
        c_tam = ws.cell(r, 17).value
        
        cod_n = norm_cod(c_cod)
        per_n = norm_periodo(c_per)
        anio_n = norm_anio(c_eval)
        
        if cod_n:
            dep_limpio = str(c_dep).strip() if c_dep else None
            tam_limpio = str(c_tam).strip() if c_tam else None
            if dep_limpio or tam_limpio:
                dict_eps_meta[(cod_n, per_n)] = {
                    'Departamento': dep_limpio,
                    'Tamano_EPS': tam_limpio
                }
        
        if cod_n and per_n and anio_n:
            key = (cod_n, per_n, anio_n)
            
            def get_cell_data(col_idx):
                cell = ws.cell(r, col_idx)
                txt = cell.value
                if txt is not None:
                    txt = str(txt).strip()
                    if txt in ('-', '', 'None', 'nan'): txt = None
                link = limpiar_link(cell.hyperlink.target if cell.hyperlink else None)
                return txt, link
            
            fisc_ini_txt, fisc_ini_lnk = get_cell_data(18)
            fisc_fin_txt, fisc_fin_lnk = get_cell_data(19)
            fisc_inf_txt, fisc_inf_lnk = get_cell_data(20)
            fisc_cmp_txt, fisc_cmp_lnk = get_cell_data(21)
            fisc_med_txt, fisc_med_lnk = get_cell_data(22)
            pas_ins_txt, pas_ins_lnk = get_cell_data(23)
            pas_cmp_txt, pas_cmp_lnk = get_cell_data(24)
            pas_dec_txt, pas_dec_lnk = get_cell_data(25)
            rec_rec_txt, rec_rec_lnk = get_cell_data(26)
            rec_ape_txt, rec_ape_lnk = get_cell_data(27)
            
            c_quinq = ws.cell(r, 7).value
            c_rango_an = ws.cell(r, 11).value
            
            def limpiar_mes_3l(m):
                return m.strip()[:3].capitalize()

            def fmt_rango_an(val):
                if not val: return None
                s = str(val).strip()
                m = re.match(r'^([a-zA-ZáéíóúÁÉÍÓÚ]{3,})\s*(\d{2,4})\s*[-–/]\s*([a-zA-ZáéíóúÁÉÍÓÚ]{3,})\s*(\d{2,4})$', s)
                if m:
                    m1, a1, m2, a2 = m.groups()
                    return f'{limpiar_mes_3l(m1)} {a1} - {limpiar_mes_3l(m2)} {a2}'
                return s

            def fmt_quinq(val):
                if not val: return None
                s = str(val).strip()
                m = re.match(r'^(\d{4})\s*[-–/]\s*(\d{4})$', s)
                if m:
                    return f'{m.group(1)} - {m.group(2)}'
                return s

            dict_inf[key] = {
                'Departamento': str(c_dep).strip() if c_dep else None,
                'Empresa_Informes': str(c_nom).strip() if c_nom else None,
                'Tamano_EPS': str(c_tam).strip() if c_tam else None,
                'Estado_Evaluacion': str(c_est).strip() if c_est else None,
                'Quinquenio': fmt_quinq(c_quinq),
                'Rango_Anio': fmt_rango_an(c_rango_an),
                'Inf_Fisc_Inicial': fisc_ini_txt,
                'Inf_Fisc_Inicial_Link': fisc_ini_lnk,
                'Inf_Fisc_Final': fisc_fin_txt,
                'Inf_Fisc_Final_Link': fisc_fin_lnk,
                'Inf_Fiscalizacion': fisc_inf_txt,
                'Inf_Fiscalizacion_Link': fisc_inf_lnk,
                'Inf_Fisc_Complementario': fisc_cmp_txt,
                'Inf_Fisc_Complementario_Link': fisc_cmp_lnk,
                'Inf_Fisc_MedidasCorrectivas': fisc_med_txt,
                'Inf_Fisc_MedidasCorrectivas_Link': fisc_med_lnk,
                'Inf_PAS_Instruccion': pas_ins_txt,
                'Inf_PAS_Instruccion_Link': pas_ins_lnk,
                'Inf_PAS_Complementario': pas_cmp_txt,
                'Inf_PAS_Complementario_Link': pas_cmp_lnk,
                'Inf_PAS_Decision': pas_dec_txt,
                'Inf_PAS_Decision_Link': pas_dec_lnk,
                'Recurso_Reconsideracion': rec_rec_txt,
                'Recurso_Reconsideracion_Link': rec_rec_lnk,
                'Recurso_Apelacion': rec_ape_txt,
                'Recurso_Apelacion_Link': rec_ape_lnk
            }
            
    print(f"Éxito: {len(dict_inf)} registros de informes cargados para cruce.")
    return dict_inf, dict_eps_meta

def parse_periodo_banner(texto, periodos_ya_vistos):
    """Detecta el periodo regulatorio en el banner de cada bloque."""
    s = texto.upper()
    if 'METAS DE SERVICIO' in s:
        return None
        
    p_detectado = None
    if 'QUINTO' in s or ' 5TO' in s or 'V PR' in s or 'V PERIODO' in s:
        p_detectado = 'V PR'
    elif 'CUARTO' in s or ' 4TO' in s or 'IV PR' in s or 'IV PERIODO' in s:
        p_detectado = 'IV PR'
    elif 'TERCER' in s or 'TERCERO' in s or ' 3ER' in s or 'III PR' in s or 'III PERIODO' in s:
        p_detectado = 'III PR'
    elif 'SEGUNDO' in s or ' 2DO' in s or 'II PR' in s or 'II PERIODO' in s:
        p_detectado = 'II PR'
    elif 'PRIMER' in s or ' 1ER' in s or 'I PR' in s or 'I PERIODO' in s:
        p_detectado = 'I PR'
        
    # Corrección de errores de copy-paste en títulos (ej: "SEGUNDO QUINQUENIO" con años <= 2010)
    m_anio = re.search(r'\((\d{4})', s)
    if m_anio:
        anio_ini = int(m_anio.group(1))
        if anio_ini <= 2010 and p_detectado in ('II PR', 'III PR'):
            p_detectado = 'I PR'
            
    # Garantizar orden lógico si ya se detectó el periodo previamente en la misma hoja
    if p_detectado in periodos_ya_vistos:
        orden = ['V PR', 'IV PR', 'III PR', 'II PR', 'I PR']
        if periodos_ya_vistos and periodos_ya_vistos[-1] in orden:
            idx = orden.index(periodos_ya_vistos[-1])
            if idx + 1 < len(orden):
                p_detectado = orden[idx + 1]

    return p_detectado

def detectar_columnas_bloque(df, r_ban, r_fin_bloque):
    """
    Detecta columnas y etapas en un bloque específico (soporta tanto bloques multi-etapa
    horizontales como tablas únicas de periodos históricos).
    """
    mapa_etapas = {}
    
    # 1. Buscar si hay etapas en filas cercanas al banner
    for num_fila in range(max(0, r_ban - 1), min(r_ban + 4, len(df))):
        fila_etapas = df.iloc[num_fila, :]
        for idx_col, valor in enumerate(fila_etapas):
            valor_str = str(valor).strip().upper()
            if 'RESULTADOS DEL CUMPLIMIENTO' in valor_str:
                continue
            if 'RESULTADO' in valor_str and 'RESULTADO_ACTUAL' not in mapa_etapas:
                mapa_etapas['RESULTADO_ACTUAL'] = idx_col
            elif 'FISC' in valor_str and 'FISCALIZACIÓN' not in mapa_etapas:
                mapa_etapas['FISCALIZACIÓN'] = idx_col
            elif 'INSTRU' in valor_str and 'INSTRUCCIÓN' not in mapa_etapas:
                mapa_etapas['INSTRUCCIÓN'] = idx_col
            elif ('DECIS' in valor_str or 'DESIC' in valor_str) and 'DECISIÓN' not in mapa_etapas:
                mapa_etapas['DECISIÓN'] = idx_col

    es_historico = len(mapa_etapas) <= 1
    if es_historico:
        # En periodos históricos la tabla única empieza en columna 5 o 6
        col_inicio_metas = 5
        mapa_etapas = {
            'FISCALIZACIÓN': col_inicio_metas,
            'RESULTADO_ACTUAL': col_inicio_metas
        }

    # 2. Buscar la fila donde están los encabezados de Años ('Año 1', 'Año 2'...)
    r_anios = None
    for r in range(r_ban + 1, min(r_ban + 6, len(df))):
        for c in range(4, min(35, len(df.columns))):
            t = str(df.iloc[r, c]).strip().upper()
            if 'AÑO 1' in t or t == '1' or 'ANO 1' in t:
                r_anios = r
                break
        if r_anios is not None:
            break
            
    if r_anios is None:
        r_anios = r_ban + 2

    # 3. Mapear columnas de métricas
    columnas_metricas = {}
    
    if es_historico:
        etapas_a_procesar = [('FISCALIZACIÓN', 4, len(df.columns)), ('RESULTADO_ACTUAL', 4, len(df.columns))]
    else:
        etapas_ordenadas = sorted(mapa_etapas.items(), key=lambda x: x[1])
        etapas_a_procesar = []
        for idx_e, (etapa, col_inicio_etapa) in enumerate(etapas_ordenadas):
            col_fin_etapa = etapas_ordenadas[idx_e + 1][1] if idx_e + 1 < len(etapas_ordenadas) else len(df.columns)
            etapas_a_procesar.append((etapa, col_inicio_etapa, col_fin_etapa))

    for etapa, col_start, col_end in etapas_a_procesar:
        if es_historico and etapa == 'RESULTADO_ACTUAL' and 'FISCALIZACIÓN' in columnas_metricas:
            columnas_metricas['RESULTADO_ACTUAL'] = columnas_metricas['FISCALIZACIÓN']
            continue
            
        metricas = {
            'Base': None,
            'Meta': {},
            'Ejecutado': {},
            'ICI': {},
            'Acumulado_Meta': None,
            'Acumulado_Ejecutado': None
        }
        
        # Buscar Base
        for c in range(max(0, col_start - 3), min(col_start + 4, len(df.columns))):
            for r in range(max(0, r_anios - 2), r_anios + 1):
                if 'BASE' in str(df.iloc[r, c] or '').upper():
                    metricas['Base'] = c
                    break
            if metricas['Base'] is not None:
                break
                
        # Detectar secciones en filas r_anios - 2 hasta r_anios
        sec_row = [None] * len(df.columns)
        for r_sec in range(max(0, r_anios - 2), r_anios):
            for c in range(col_start, col_end):
                val = str(df.iloc[r_sec, c] or '').strip().upper()
                if 'META' in val and 'BASE' not in val and 'UNIDAD' not in val:
                    sec_row[c] = 'Meta'
                elif 'EJEC' in val or 'OBTENIDO' in val:
                    sec_row[c] = 'Ejecutado'
                elif 'ICI' in val or 'CUMPLIMIENTO' in val:
                    sec_row[c] = 'ICI'
                    
        has_sec = any(sec_row[c] is not None for c in range(col_start, col_end))
        
        if has_sec:
            curr_sec = None
            for c in range(col_start, col_end):
                if sec_row[c]: curr_sec = sec_row[c]
                else: sec_row[c] = curr_sec
                
            for c in range(col_start, col_end):
                sec = sec_row[c]
                if not sec: continue
                t = str(df.iloc[r_anios, c] or '').strip().upper()
                if not t or t == 'NAN':
                    if r_anios + 1 < len(df):
                        t = str(df.iloc[r_anios + 1, c] or '').strip().upper()
                if not t or t == 'NAN': continue
                
                m_pt = re.search(r'PT\s*(\d+)', t)
                m_a = re.search(r'(?:AÑO|ANO)?\s*([1-9])\b', t)
                if 'ACUM' in t or 'ACUM' in str(df.iloc[r_anios - 1, c] or '').upper():
                    if sec == 'Meta' and metricas['Acumulado_Meta'] is None:
                        metricas['Acumulado_Meta'] = c
                    elif sec in ('Ejecutado', 'ICI') and metricas['Acumulado_Ejecutado'] is None:
                        metricas['Acumulado_Ejecutado'] = c
                elif m_pt:
                    metricas[sec][f'PT {m_pt.group(1)}'] = c
                elif m_a:
                    metricas[sec][int(m_a.group(1))] = c
        else:
            # Fallback counting
            contadores_anio = {i: 0 for i in range(1, 10)}
            contadores_pt = {}
            for c in range(col_start, col_end):
                t = str(df.iloc[r_anios, c] or '').strip().upper()
                if not t or t == 'NAN':
                    if r_anios + 1 < len(df):
                        t = str(df.iloc[r_anios + 1, c] or '').strip().upper()
                if not t or t == 'NAN': continue
                
                m_pt = re.search(r'PT\s*(\d+)', t)
                m_a = re.search(r'(?:AÑO|ANO)?\s*([1-9])\b', t)
                if 'ACUM' in t or 'ACUM' in str(df.iloc[r_anios - 1, c] or '').upper():
                    if metricas['Acumulado_Meta'] is None:
                        metricas['Acumulado_Meta'] = c
                    elif metricas['Acumulado_Ejecutado'] is None:
                        metricas['Acumulado_Ejecutado'] = c
                elif m_pt:
                    pt_k = f"PT {m_pt.group(1)}"
                    contadores_pt[pt_k] = contadores_pt.get(pt_k, 0) + 1
                    ap = contadores_pt[pt_k]
                    if ap == 1: metricas['Ejecutado'][pt_k] = c
                    elif ap == 2: metricas['ICI'][pt_k] = c
                elif m_a:
                    a_num = int(m_a.group(1))
                    contadores_anio[a_num] += 1
                    ap = contadores_anio[a_num]
                    if ap == 1: metricas['Meta'][a_num] = c
                    elif ap == 2: metricas['Ejecutado'][a_num] = c
                    elif ap == 3: metricas['ICI'][a_num] = c

        columnas_metricas[etapa] = metricas

    # Determinar fila de inicio de metas
    r_start_metas = r_anios + 1
    for r_check in range(r_anios + 1, min(r_anios + 4, len(df))):
        c2 = str(df.iloc[r_check, 1]).strip()
        c3 = str(df.iloc[r_check, 2]).strip()
        if c2.isdigit() or (c3 and c3.lower() not in ('nan', '') and len(c3) > 3 and not any(k in c3.upper() for k in ['AÑO', 'META', 'UNIDAD'])):
            r_start_metas = r_check
            break
            
    return mapa_etapas, columnas_metricas, r_start_metas

def extraer_todas_las_eps_dinamico():
    file_path = "Compilatorio ICI_ICG_EPS TOTALES_ACT EN PROCESO.xlsx"
    print("Iniciando extracción multi-periodo nacional (todos los periodos regulatorios por EPS)...")
    
    # 1. Cargar metadatos de informes linkeados
    dict_inf, dict_eps_meta = cargar_informes_totales(file_path)
    
    try:
        xls = pd.ExcelFile(file_path)
    except Exception as e:
        print(f"Error al cargar el archivo Excel: {e}")
        return
        
    # Incluir todas las hojas numéricas Y la EPS 008-052 (ATUSA)
    hojas_eps = [hoja for hoja in xls.sheet_names if hoja.isdigit() or hoja == '008-052']
    print(f"Total de hojas EPS a procesar: {len(hojas_eps)}")
    
    datos_extraidos_totales = []
    total_bloques_procesados = 0
    
    for hoja in hojas_eps:
        df = pd.read_excel(xls, sheet_name=hoja, header=None)
        
        cod_ep = hoja
        nombre_eps = str(df.iloc[1, 0]).strip() if pd.notna(df.iloc[1, 0]) else f"EPS {hoja}"
        cod_ep_num = norm_cod(cod_ep)
        
        # 1. Encontrar todos los bloques de periodos regulatorios en la hoja
        bloques = []
        vistos = []
        for r in range(len(df)):
            for c in range(min(5, len(df.columns))):
                val = str(df.iloc[r, c] or '').strip().replace('\n', ' ')
                if "RESULTADOS DEL CUMPLIMIENTO DE METAS DE GESTI" in val.upper():
                    pr = parse_periodo_banner(val, vistos)
                    if pr:
                        vistos.append(pr)
                        bloques.append((r, pr))
                    break
                    
        # Fallback si no se detectó ningún banner: usar celdas fijas tradicionales
        if not bloques:
            pr_fallback = str(df.iloc[2, 0]).strip() if pd.notna(df.iloc[2, 0]) else "III PR"
            bloques = [(3, pr_fallback)]
            
        total_bloques_procesados += len(bloques)
        
        # 2. Procesar cada bloque (periodo regulatorio) de la EPS
        for idx_b, (r_ban, periodo_reg) in enumerate(bloques):
            r_fin_bloque = bloques[idx_b + 1][0] if idx_b + 1 < len(bloques) else len(df)
            periodo_reg_romano = norm_periodo(periodo_reg)
            
            mapa_etapas, columnas_metricas, r_start = detectar_columnas_bloque(df, r_ban, r_fin_bloque)
            
            meta_eps_gen = dict_eps_meta.get((cod_ep_num, periodo_reg_romano), {})
            departamento_val = meta_eps_gen.get('Departamento')
            tamano_val = meta_eps_gen.get('Tamano_EPS')
            
            memoria_id_meta = ""
            memoria_meta_principal = ""
            
            for row_idx in range(r_start, r_fin_bloque):
                id_meta_crudo = str(df.iloc[row_idx, 1]).strip()
                desc_meta_crudo = str(df.iloc[row_idx, 2]).strip()
                unidad_medida_cruda = str(df.iloc[row_idx, 3]).strip()
                if unidad_medida_cruda.lower() in ('nan', '-'): unidad_medida_cruda = None
                
                texto_combinado = (id_meta_crudo + " " + desc_meta_crudo).lower()
                
                es_fila_icg = False
                if 'índice de cumplimiento global' in texto_combinado or 'indice de cumplimiento global' in texto_combinado or 'icg' in texto_combinado:
                    es_fila_icg = True
                elif any(k in texto_combinado for k in ['fuente:', 'elaboración:', 'elaboracion:', 'nota:']):
                    break
                    
                if (id_meta_crudo.lower() in ('nan', '')) and (desc_meta_crudo.lower() in ('nan', '')) and not es_fila_icg:
                    continue
                    
                if es_fila_icg:
                    nivel = "EPS"
                    id_meta_crudo = "ICG"
                    memoria_meta_principal = "Índice de Cumplimiento Global (ICG)"
                    desc_especifica = "Índice de Cumplimiento Global (ICG)"
                    unidad_medida_cruda = "%"
                else:
                    if id_meta_crudo not in ('nan', ''):
                        nivel = "EPS"
                        memoria_id_meta = id_meta_crudo
                        memoria_meta_principal = desc_meta_crudo
                        desc_especifica = desc_meta_crudo
                    else:
                        nivel = "Localidad"
                        id_meta_crudo = memoria_id_meta
                        desc_especifica = desc_meta_crudo
                        
                for etapa in mapa_etapas.keys():
                    metricas = columnas_metricas.get(etapa)
                    if not metricas: continue
                    
                    def agregar_registro(anio_etiqueta, v_meta, v_ejec, v_ici):
                        if v_meta is None and v_ejec is None and v_ici is None:
                            return
                        info_cruce = dict_inf.get((cod_ep_num, periodo_reg_romano, anio_etiqueta), {})
                        
                        inf_etapa = None
                        inf_etapa_link = None
                        if etapa == 'FISCALIZACIÓN':
                            inf_etapa = info_cruce.get('Inf_Fiscalizacion') or info_cruce.get('Inf_Fisc_Final') or info_cruce.get('Inf_Fisc_Inicial')
                            inf_etapa_link = info_cruce.get('Inf_Fiscalizacion_Link') or info_cruce.get('Inf_Fisc_Final_Link') or info_cruce.get('Inf_Fisc_Inicial_Link')
                        elif etapa == 'INSTRUCCIÓN':
                            inf_etapa = info_cruce.get('Inf_PAS_Instruccion')
                            inf_etapa_link = info_cruce.get('Inf_PAS_Instruccion_Link')
                        elif etapa == 'DECISIÓN':
                            inf_etapa = info_cruce.get('Inf_PAS_Decision')
                            inf_etapa_link = info_cruce.get('Inf_PAS_Decision_Link')
                        elif etapa == 'RESULTADO_ACTUAL':
                            inf_etapa = info_cruce.get('Inf_Fiscalizacion') or info_cruce.get('Inf_Fisc_Final') or info_cruce.get('Inf_Fisc_Inicial')
                            inf_etapa_link = info_cruce.get('Inf_Fiscalizacion_Link') or info_cruce.get('Inf_Fisc_Final_Link') or info_cruce.get('Inf_Fisc_Inicial_Link')
                            
                        datos_extraidos_totales.append({
                            'Cod_EP': cod_ep,
                            'Nombre_EPS': nombre_eps,
                            'Periodo_Regulatorio': periodo_reg,
                            'Departamento': info_cruce.get('Departamento') or departamento_val,
                            'Tamano_EPS': info_cruce.get('Tamano_EPS') or tamano_val,
                            'Estado_Evaluacion': info_cruce.get('Estado_Evaluacion'),
                            'Quinquenio': info_cruce.get('Quinquenio'),
                            'Rango_Anio': info_cruce.get('Rango_Anio'),
                            'ID_Meta_General': id_meta_crudo,
                            'Meta_Principal': memoria_meta_principal,
                            'Nivel_Evaluacion': nivel,
                            'Descripcion_Especifica': desc_especifica,
                            'Unidad_Medida': unidad_medida_cruda,
                            'Etapa': etapa,
                            'Año_Regulatorio': anio_etiqueta,
                            'Valor_Exigido': v_meta,
                            'Valor_Ejecutado': v_ejec,
                            'ICI': v_ici,
                            'Informe_Etapa': inf_etapa,
                            'Informe_Etapa_Link': inf_etapa_link,
                            'Inf_Fisc_Inicial': info_cruce.get('Inf_Fisc_Inicial'),
                            'Inf_Fisc_Inicial_Link': info_cruce.get('Inf_Fisc_Inicial_Link'),
                            'Inf_Fisc_Final': info_cruce.get('Inf_Fisc_Final'),
                            'Inf_Fisc_Final_Link': info_cruce.get('Inf_Fisc_Final_Link'),
                            'Inf_Fiscalizacion': info_cruce.get('Inf_Fiscalizacion'),
                            'Inf_Fiscalizacion_Link': info_cruce.get('Inf_Fiscalizacion_Link'),
                            'Inf_Fisc_Complementario': info_cruce.get('Inf_Fisc_Complementario'),
                            'Inf_Fisc_Complementario_Link': info_cruce.get('Inf_Fisc_Complementario_Link'),
                            'Inf_Fisc_MedidasCorrectivas': info_cruce.get('Inf_Fisc_MedidasCorrectivas'),
                            'Inf_Fisc_MedidasCorrectivas_Link': info_cruce.get('Inf_Fisc_MedidasCorrectivas_Link'),
                            'Inf_PAS_Instruccion': info_cruce.get('Inf_PAS_Instruccion'),
                            'Inf_PAS_Instruccion_Link': info_cruce.get('Inf_PAS_Instruccion_Link'),
                            'Inf_PAS_Complementario': info_cruce.get('Inf_PAS_Complementario'),
                            'Inf_PAS_Complementario_Link': info_cruce.get('Inf_PAS_Complementario_Link'),
                            'Inf_PAS_Decision': info_cruce.get('Inf_PAS_Decision'),
                            'Inf_PAS_Decision_Link': info_cruce.get('Inf_PAS_Decision_Link'),
                            'Recurso_Reconsideracion': info_cruce.get('Recurso_Reconsideracion'),
                            'Recurso_Reconsideracion_Link': info_cruce.get('Recurso_Reconsideracion_Link'),
                            'Recurso_Apelacion': info_cruce.get('Recurso_Apelacion'),
                            'Recurso_Apelacion_Link': info_cruce.get('Recurso_Apelacion_Link')
                        })
                        
                    # 1. AÑO BASE
                    if not es_fila_icg and metricas['Base'] is not None:
                        val_base = limpiar_valor(df.iloc[row_idx, metricas['Base']], unidad_medida_cruda)
                        agregar_registro('Año Base', val_base, None, None)
                        
                    # 2. AÑOS REGULATORIOS (1..5) Y AÑOS TRANSITORIOS (PT 1..N)
                    anios_nums = set()
                    for s in ('Meta', 'Ejecutado', 'ICI'):
                        for k in metricas[s].keys():
                            if isinstance(k, int):
                                anios_nums.add(k)
                    anios_ordenados = [f"Año {a}" for a in sorted(anios_nums)]
                    
                    pts_encontrados = set()
                    for s in ('Meta', 'Ejecutado', 'ICI'):
                        for k in metricas[s].keys():
                            if isinstance(k, str) and k.startswith('PT'):
                                pts_encontrados.add(k)
                                
                    def orden_pt(pt_str):
                        m = re.search(r'\d+', pt_str)
                        return int(m.group()) if m else 99
                        
                    for pt in sorted(pts_encontrados, key=orden_pt):
                        anios_ordenados.append(pt)
                        
                    for etiqueta_anio in anios_ordenados:
                        if etiqueta_anio.startswith("Año "):
                            clave_k = int(etiqueta_anio.replace("Año ", ""))
                        else:
                            clave_k = etiqueta_anio
                            
                        idx_meta = metricas['Meta'].get(clave_k)
                        idx_ejec = metricas['Ejecutado'].get(clave_k)
                        idx_ici = metricas['ICI'].get(clave_k)
                        
                        if es_fila_icg:
                            val_meta = None
                            val_ejec = None
                            val_ici = limpiar_valor(df.iloc[row_idx, idx_ici], '%', es_ici=True) if idx_ici else None
                        else:
                            val_meta = limpiar_valor(df.iloc[row_idx, idx_meta], unidad_medida_cruda) if idx_meta else None
                            val_ejec = limpiar_valor(df.iloc[row_idx, idx_ejec], unidad_medida_cruda) if idx_ejec else None
                            val_ici = limpiar_valor(df.iloc[row_idx, idx_ici], unidad_medida_cruda, es_ici=True) if idx_ici else None
                            
                        agregar_registro(etiqueta_anio, val_meta, val_ejec, val_ici)
                        
                    # 3. ACUMULADO
                    if not es_fila_icg:
                        idx_acum_m = metricas['Acumulado_Meta']
                        idx_acum_e = metricas['Acumulado_Ejecutado']
                        val_acum_m = limpiar_valor(df.iloc[row_idx, idx_acum_m], unidad_medida_cruda) if idx_acum_m else None
                        val_acum_e = limpiar_valor(df.iloc[row_idx, idx_acum_e], unidad_medida_cruda) if idx_acum_e else None
                        agregar_registro('Acumulado', val_acum_m, val_acum_e, None)
                        
                if es_fila_icg:
                    break

    if datos_extraidos_totales:
        df_final = pd.DataFrame(datos_extraidos_totales)
        output_name = "Base_EPS_Nacional_Estructurada_Final.xlsx"
        print(f"Escribiendo resultado en '{output_name}'...")
        df_final.to_excel(output_name, index=False)
        print(f"¡Éxito total! Se extrajeron {len(df_final)} registros enriquecidos.")
        print(f"EPS procesadas: {df_final['Cod_EP'].nunique()}")
        print(f"Total bloques/periodos procesados: {total_bloques_procesados}")
        print(f"Periodos regulatorios encontrados: {df_final['Periodo_Regulatorio'].unique()}")
        print(f"Años presentes: {df_final['Año_Regulatorio'].unique()}")
        print(f"Registros con informe vinculado: {df_final['Informe_Etapa'].notna().sum()}")

if __name__ == "__main__":
    extraer_todas_las_eps_dinamico()